"""Build the KSRI seminar deck (2026-08, v2 after Thanos's feedback) in the KIT template.

v2 changes (feedback 2026-08-08): icon-driven motivation with the practical
problem; categorised background (what / how / examples); FI slide that defines
FI first, then FI_in and FI_out (derivation verified against
Section_7_Functional_Information_for_Prompts.md §7.3/§7.4); gap slide that
derives RQ1-RQ3; icon pipelines for design + instrument; native editable CHARTS
instead of tables on all results slides (tables moved to backup); icon-flow
probe slide; contributions mapped to RQs + a unified RQ/answer slide;
limitations & outlook; key takeaways; future work. Style: body text >= 14 pt,
chevron take-home lines, no em dashes, icons built from native shapes.

Every number is frozen and verified; sources:
  [SH]  data/stats_hygiene.md            (endpoints, CIs, Holm/BH, pooled, reliability)
  [R1]  RESULTS_R1_R2_R3_2026-08-07.md   (lottery, hierarchical rho_F, bounds)
  [R6]  data/width_dial_analysis.md + RESULTS_R6_width_dial_2026-08-08.md
  [MR]  data/metric_reductions.md        (identities, POSIX/Williams)
  [PE]  data/probe_eval_hardened.md      (probe numbers)
  [CV]  RHO_F_CONSTRUCT_VALIDITY_2026-08-07.md (payoff prediction)
  [FA]  factor_structure on the canonical matrix (74.6 %, Horn 3)

    uv run python -m prompt_sensitivity.scripts.make_ksri_deck
"""

from __future__ import annotations

import argparse
import shutil

from loguru import logger
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION,
                             XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ..config import load_config
from ..logging_setup import configure_logging

TEMPLATE = "Prompt_Sensitivity_Full_Results_final.pptx"
OUT = "KSRI_Seminar_Prompt_Sensitivity_2026-08.pptx"

LAYOUT_TITLE_TEXT = 13
LAYOUT_AGENDA = 6
LAYOUT_DIVIDER = 11

KIT_GREEN = RGBColor(0x00, 0x96, 0x82)
KIT_BLUE = RGBColor(0x46, 0x64, 0xAA)
KIT_GREEN_TINT = RGBColor(0xD9, 0xEF, 0xEB)
KIT_BLUE_TINT = RGBColor(0xE0, 0xE6, 0xF2)
KIT_ORANGE = RGBColor(0xDF, 0xA0, 0x1D)
KIT_RED = RGBColor(0xA2, 0x22, 0x23)
RED_TINT = RGBColor(0xF6, 0xE1, 0xE1)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF2, 0xF2, 0xF2)
GRAY = RGBColor(0x9B, 0x9B, 0x9B)
FONT = "Arial"

# consistent model colours across every chart
C_QWEN, C_LLAMA, C_MISTRAL = KIT_GREEN, KIT_BLUE, KIT_ORANGE

FUNNEL_SHAPE = getattr(MSO_SHAPE, "FUNNEL", MSO_SHAPE.FLOWCHART_MERGE)
GEAR_SHAPE = getattr(MSO_SHAPE, "GEAR_6", MSO_SHAPE.OVAL)


def In(v: float):
    return Inches(v)


# --------------------------------------------------------------------------- #
# low-level helpers                                                           #
# --------------------------------------------------------------------------- #


def drop_slides_from(prs, first_1based: int) -> int:
    xml_slides = prs.slides._sldIdLst
    removed = 0
    for sld in list(xml_slides)[first_1based - 1:]:
        rId = sld.get(qn("r:id"))
        prs.part.drop_rel(rId)
        xml_slides.remove(sld)
        removed += 1
    return removed


def _style(run, size, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT


def _fill_tf(tf, lines, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = [(line, {})] if isinstance(line, str) else [
            (rn, {}) if isinstance(rn, str) else rn for rn in line]
        for text, ov in runs:
            r = p.add_run()
            r.text = text
            _style(r, ov.get("size", size), ov.get("bold", bold),
                   ov.get("color", color), ov.get("italic", False))


def T(slide, x, y, w, h, lines, size=14, color=INK, bold=False,
      align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    _fill_tf(tb.text_frame, lines, size=size, color=color, bold=bold, align=align)
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, rot=0):
    sp = slide.shapes.add_shape(kind, In(x), In(y), In(w), In(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if rot:
        sp.rotation = rot
    return sp


def box(slide, x, y, w, h, lines, fill=KIT_GREEN_TINT, line_color=KIT_GREEN,
        size=14, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = shape(slide, kind, x, y, w, h, fill=fill, line=line_color)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.07
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = In(0.10)
    tf.margin_top = tf.margin_bottom = In(0.05)
    tf.vertical_anchor = anchor
    _fill_tf(tf, lines, size=size, color=color, align=align)
    return sp


def arrow(slide, x1, y1, x2, y2, color=MUTED, width_pt=1.75):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      In(x1), In(y1), In(x2), In(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width_pt)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "med", "len": "med"}))
    return conn


def chevron(slide, text, top=6.25, size=15):
    """Take-home line: small KIT-green chevron + bold green text."""
    shape(slide, MSO_SHAPE.CHEVRON, 0.41, top + 0.04, 0.26, 0.20,
          fill=KIT_GREEN, line=None)
    return T(slide, 0.76, top, 12.2, 0.5, [text], size=size, color=KIT_GREEN,
             bold=True)


def kit_slide(prs, title, subtitle=None, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_TEXT])
    slide.shapes.title.text_frame.text = title
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = FONT
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 13:
            ph._element.getparent().remove(ph._element)
    if subtitle:
        T(slide, 0.41, 1.16, 12.5, 0.4, [subtitle], size=14, color=MUTED)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def divider(prs, title, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_DIVIDER])
    slide.shapes.title.text_frame.text = title
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = FONT
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in (13, 14):
            ph._element.getparent().remove(ph._element)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_table(slide, x, y, w, data, col_widths=None, row_height=0.32,
              header_fill=KIT_BLUE, size=12, header_size=12,
              first_col_bold=False, aligns=None, highlight_rows=(),
              highlight_fill=KIT_GREEN_TINT):
    rows, cols = len(data), len(data[0])
    gf = slide.shapes.add_table(rows, cols, In(x), In(y), In(w),
                                In(row_height * rows))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    if col_widths:
        total = sum(col_widths)
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Emu(int(In(w) * cw / total))
    for i, row in enumerate(data):
        tbl.rows[i].height = In(row_height)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = In(0.06)
            cell.margin_top = cell.margin_bottom = In(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif i in highlight_rows:
                cell.fill.fore_color.rgb = highlight_fill
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else ROW_ALT
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (aligns[j] if aligns else
                           (PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER))
            for k, seg in enumerate(str(val).split("\n")):
                pp = p if k == 0 else tf.add_paragraph()
                if k > 0:
                    pp.alignment = p.alignment
                r = pp.add_run()
                r.text = seg
                _style(r, header_size if i == 0 else size,
                       bold=(i == 0) or (j == 0 and first_col_bold),
                       color=WHITE if i == 0 else INK)
    return gf


def add_chart(slide, x, y, w, h, ctype, cats, series, *, colors, title=None,
              y_min=None, y_max=None, fmt="0.00", labels=True, legend=True,
              label_size=10, major_unit=None, label_fmt=None):
    """Native, fully editable PowerPoint chart in KIT colours."""
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    gf = slide.shapes.add_chart(ctype, In(x), In(y), In(w), In(h), data)
    ch = gf.chart
    ch.has_title = title is not None
    if title is not None:
        tf = ch.chart_title.text_frame
        tf.text = title
        for r in tf.paragraphs[0].runs:
            _style(r, 13, bold=True)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
        ch.legend.font.name = FONT
    va = ch.value_axis
    if y_min is not None:
        va.minimum_scale = y_min
    if y_max is not None:
        va.maximum_scale = y_max
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.name = FONT
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    if major_unit is not None:
        va.major_unit = major_unit
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.name = FONT
    if y_min is not None and y_min < 0:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    line_chart = ctype in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS)
    for i, s in enumerate(ch.series):
        col = colors[i % len(colors)]
        if line_chart:
            s.format.line.color.rgb = col
            s.format.line.width = Pt(2.75)
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = col
            s.format.line.fill.background()
            ser = s._element
            inv = ser.find(qn("c:invertIfNegative"))
            if inv is None:
                inv = ser.makeelement(qn("c:invertIfNegative"), {})
                spPr = ser.find(qn("c:spPr"))
                spPr.addnext(inv)
            inv.set("val", "0")
    if labels and not line_chart:
        plot = ch.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = label_fmt or fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(label_size)
        dl.font.name = FONT
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    return gf


def _retext(tf, lines, keep_size=True):
    proto = None
    for p in tf.paragraphs:
        if p.runs:
            proto = p.runs[0].font
            break
    kw = {}
    if proto is not None:
        kw = dict(size=(proto.size.pt if (keep_size and proto.size) else 18),
                  bold=bool(proto.bold), name=proto.name or FONT)
        try:
            kw["color"] = proto.color.rgb
        except (AttributeError, TypeError):
            kw["color"] = None
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        if kw:
            r.font.size = Pt(kw["size"])
            r.font.bold = kw["bold"]
            r.font.name = kw["name"]
            if kw.get("color") is not None:
                r.font.color.rgb = kw["color"]


# --------------------------------------------------------------------------- #
# icons (native shapes only, all editable)                                    #
# --------------------------------------------------------------------------- #


def icon_person(s, cx, y, size=0.72, color=KIT_BLUE):
    shape(s, MSO_SHAPE.OVAL, cx - size * 0.18, y, size * 0.36, size * 0.36,
          fill=color, line=None)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - size * 0.32, y + size * 0.40,
          size * 0.64, size * 0.50, fill=color, line=None)


def icon_bubble(s, x, y, w, h, text, fill=KIT_BLUE_TINT, line=KIT_BLUE,
                size=13):
    sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, x, y, w, h,
               fill=fill, line=line)
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = In(0.07)
    tf.margin_top = tf.margin_bottom = In(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    _fill_tf(tf, [text], size=size, align=PP_ALIGN.CENTER)
    return sp


def icon_outcome(s, cx, y, ok, d=0.42):
    sp = shape(s, MSO_SHAPE.OVAL, cx - d / 2, y, d, d,
               fill=KIT_GREEN_TINT if ok else RED_TINT,
               line=KIT_GREEN if ok else KIT_RED, line_w=1.5)
    tf = sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_tf(tf, [[("✓" if ok else "✗",
                    {"bold": True, "size": 16,
                     "color": KIT_GREEN if ok else KIT_RED})]],
             align=PP_ALIGN.CENTER)
    return sp


def icon_station(s, cx, y, kind, label, sub=None, fill=KIT_BLUE_TINT,
                 line=KIT_BLUE, ih=0.85, iw=0.95, label_w=2.0):
    """Icon shape centred at cx with a bold label and optional caption below."""
    if kind == "can":
        shape(s, MSO_SHAPE.CAN, cx - iw / 2 + 0.12, y, iw - 0.24, ih,
              fill=fill, line=line)
    elif kind == "funnel":
        shape(s, FUNNEL_SHAPE, cx - iw / 2 + 0.08, y, iw - 0.16, ih,
              fill=fill, line=line)
    elif kind == "gear":
        shape(s, GEAR_SHAPE, cx - ih / 2, y, ih, ih, fill=fill, line=line)
    elif kind == "cube":
        shape(s, MSO_SHAPE.CUBE, cx - iw / 2 + 0.06, y + 0.04, iw - 0.12,
              ih - 0.08, fill=fill, line=line)
    elif kind == "cubes":
        for k in range(3):
            shape(s, MSO_SHAPE.CUBE, cx - 0.52 + k * 0.22, y + 0.24 - k * 0.11,
                  0.62, 0.55, fill=fill, line=line)
    elif kind == "check":
        icon_outcome(s, cx, y + 0.16, True, d=0.55)
    elif kind == "levels":
        icon_bubble(s, cx - 0.62, y, 0.62, 0.38, "L0", size=12,
                    fill=RED_TINT, line=KIT_RED)
        icon_bubble(s, cx + 0.02, y + 0.42, 0.62, 0.38, "L1", size=12,
                    fill=KIT_GREEN_TINT, line=KIT_GREEN)
    elif kind == "grid":
        for k in range(10):
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE,
                  cx - 0.52 + (k % 5) * 0.22, y + 0.14 + (k // 5) * 0.30,
                  0.18, 0.24, fill=KIT_GREEN_TINT, line=KIT_GREEN)
    elif kind == "dedup":
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.34, y + 0.10, 0.5, 0.55,
              fill=ROW_ALT, line=GRAY)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.14, y + 0.22, 0.5, 0.55,
              fill=WHITE, line=KIT_BLUE)
    elif kind == "bubble":
        sp = icon_bubble(s, cx - 0.45, y + 0.06, 0.9, 0.62, "?", size=20)
        for p in sp.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
    elif kind == "slabs":
        for k in range(3):
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.42 + k * 0.07,
                  y + 0.10 + k * 0.20, 0.78, 0.15, fill=KIT_BLUE_TINT,
                  line=KIT_BLUE)
    elif kind == "blockarrow":
        shape(s, MSO_SHAPE.RIGHT_ARROW, cx - 0.45, y + 0.18, 0.9, 0.44,
              fill=KIT_GREEN_TINT, line=KIT_GREEN)
    elif kind == "head":
        sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.42, y + 0.16, 0.84,
                   0.5, fill=WHITE, line=KIT_BLUE)
        tf = sp.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _fill_tf(tf, [[("w·x", {"size": 14, "bold": True,
                                     "color": KIT_BLUE})]],
                 align=PP_ALIGN.CENTER)
    elif kind == "warn":
        sp = shape(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.38, y + 0.04, 0.76,
                   0.66, fill=RGBColor(0xFA, 0xF0, 0xDC), line=KIT_ORANGE,
                   line_w=1.5)
        tf = sp.text_frame
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        _fill_tf(tf, [[("!", {"size": 15, "bold": True,
                              "color": KIT_ORANGE})]], align=PP_ALIGN.CENTER)
    elif kind == "scatter":
        for k, (dx, dy, dd) in enumerate([(0.0, 0.05, 0.16), (0.35, 0.22, 0.2),
                                          (-0.3, 0.3, 0.14), (0.1, 0.48, 0.18),
                                          (-0.15, 0.6, 0.13)]):
            shape(s, MSO_SHAPE.OVAL, cx + dx - dd / 2, y + dy, dd, dd,
                  fill=KIT_BLUE, line=None)
    elif kind == "barsdown":
        for k, hh in enumerate([0.62, 0.44, 0.26]):
            shape(s, MSO_SHAPE.RECTANGLE, cx - 0.36 + k * 0.28,
                  y + 0.72 - hh, 0.2, hh, fill=KIT_BLUE, line=None)
    elif kind == "pie":
        shape(s, MSO_SHAPE.PIE, cx - 0.36, y + 0.04, 0.72, 0.72,
              fill=KIT_BLUE_TINT, line=KIT_BLUE)
    elif kind == "cycle":
        shape(s, MSO_SHAPE.CIRCULAR_ARROW, cx - 0.38, y + 0.02, 0.76, 0.72,
              fill=KIT_BLUE_TINT, line=KIT_BLUE)
    elif kind == "lock":
        shape(s, MSO_SHAPE.ARC, cx - 0.20, y + 0.02, 0.40, 0.36,
              fill=None, line=KIT_RED, line_w=2.5)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.28, y + 0.24, 0.56, 0.42,
              fill=RED_TINT, line=KIT_RED)
    elif kind == "dial":
        shape(s, MSO_SHAPE.BLOCK_ARC, cx - 0.40, y + 0.08, 0.80, 0.70,
              fill=KIT_GREEN_TINT, line=KIT_GREEN)
        arrow(s, cx, y + 0.62, cx + 0.24, y + 0.24, color=KIT_GREEN,
              width_pt=2.0)
    if label:
        T(s, cx - label_w / 2, y + 0.92, label_w, 0.3, [label], size=14,
          bold=True, align=PP_ALIGN.CENTER)
    if sub:
        T(s, cx - label_w / 2, y + 1.20, label_w, 0.55, [sub], size=12,
          color=MUTED, align=PP_ALIGN.CENTER)


# --------------------------------------------------------------------------- #
# main-path slides                                                            #
# --------------------------------------------------------------------------- #


def s_title(prs):
    s = prs.slides[0]
    for sh in s.shapes:
        if sh.name == "Titel 4":
            _retext(sh.text_frame,
                    ["Prompt Sensitivity Is Three", "Measurements, Not One"])
        elif sh.name == "Textfeld 17":
            _retext(sh.text_frame,
                    ["A measurement model for how LLMs respond to the phrasing of a question",
                     "Athanasios Drossos  ·  seminar presentation  ·  KSRI / KIT  ·  August 2026"],
                    keep_size=False)
        elif sh.name == "Datumsplatzhalter 18":
            _retext(sh.text_frame, ["August 2026"])
        elif sh.name == "Fußzeilenplatzhalter 19":
            _retext(sh.text_frame,
                    ["Drossos · Prompt Sensitivity Is Three Measurements"])
    s.notes_slide.notes_text_frame.text = (
        "~30 s. One-liner: 'prompt sensitivity' is treated as one number; we show it is three "
        "separate measurements, give the missing one an estimator, and validate each axis with "
        "its own intervention.")


def s_agenda(prs):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_AGENDA])
    s.shapes.title.text_frame.text = "Agenda"
    body = next(ph for ph in s.placeholders if ph.placeholder_format.idx == 13)
    items = [
        "Motivation",
        "Background: metrics, functional information, the AmbigQA testbed",
        "Gap and research question",
        "Methods: design, pipeline, the three-axis model",
        "Results: two dials, one double dissociation",
        "The prompt checker",
        "Contributions, limitations, takeaways, future work",
    ]
    tf = body.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = it
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(17)
    s.notes_slide.notes_text_frame.text = "~20 s. Structure mirrors a paper."


def s_motivation(prs):
    s = kit_slide(prs, "One question, three phrasings, three outcomes", None,
                  notes="~90 s. Walk the picture left to right: a user picks ONE of many equivalent "
                        "phrasings without knowing it matters; the model answers correctly for some "
                        "phrasings and not others. Then the three consequences. Pattern is "
                        "illustrative; the phenomenon is real and quantified later.")
    # user
    icon_person(s, 1.0, 1.85)
    T(s, 0.35, 2.70, 1.3, 0.3, ["a user"], size=14, bold=True,
      align=PP_ALIGN.CENTER)
    T(s, 0.15, 3.00, 1.7, 0.6, ["picks one phrasing,", "unaware it matters"],
      size=12, color=MUTED, align=PP_ALIGN.CENTER)
    # three equivalent phrasings
    bubbles = ["“When did The Nun come out?”",
               "“What year was The Nun released?”",
               "“The Nun: release date?”"]
    for i, b in enumerate(bubbles):
        icon_bubble(s, 2.15, 1.55 + i * 0.95, 3.55, 0.72, b, size=13)
        arrow(s, 1.55, 2.15 + (i - 1) * 0.22, 2.12, 1.92 + i * 0.95)
    T(s, 2.15, 4.35, 3.55, 0.3, ["same meaning, verified equivalent"],
      size=12, color=MUTED, align=PP_ALIGN.CENTER)
    # LLM
    icon_station(s, 6.75, 1.95, "cube", "LLM", None, iw=1.15, ih=1.0)
    for i in range(3):
        arrow(s, 5.75, 1.91 + i * 0.95, 6.15, 2.30 + (i - 1) * 0.28)
        arrow(s, 7.40, 2.30 + (i - 1) * 0.28, 7.90, 1.86 + i * 0.95)
    # outcomes
    for i, ok in enumerate([True, False, False]):
        icon_outcome(s, 8.20, 1.68 + i * 0.95, ok)
    # right: the practical problem
    T(s, 9.05, 1.50, 3.9, 0.9,
      ["The user cannot know, before asking,", "which phrasing will succeed."],
      size=16, bold=True)
    T(s, 9.05, 2.45, 3.9, 1.9,
      ["No signal marks the losing phrasings.",
       "Prompt advice is folklore, not measurement.",
       "Reliability of deployed LLM services", "varies with wording, unmanaged."],
      size=14)
    # consequences
    for cx, kind, txt in [
            (1.55, "cycle", "prompt engineering stays\ntrial and error"),
            (5.35, "barsdown", "suboptimal phrasing\nsilently costs accuracy"),
            (9.15, "warn", "per-question reliability\nis unmeasured")]:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.65, 4.78, 3.7, 1.08,
              fill=WHITE, line=GRAY)
        icon_station(s, cx - 0.05, 4.92, kind, None)
        T(s, cx + 0.55, 4.98, 2.42, 0.75, txt.split("\n"), size=14)
    chevron(s, "Phrasing, not knowledge, often decides the outcome. This talk: measure that, per question.")


def s_background(prs):
    s = kit_slide(prs, "Background: three families of measures", None,
                  notes="~90 s. Organise the literature by WHAT is measured and HOW, then name examples. "
                        "Punchline stays a teaser: Result 1 proves the first family is largely one quantity.")
    cards = [
        ("scatter", "Answer variation",
         "How much answers differ across\nphrasings or samples.",
         "Sample or rephrase, then compare\nthe answers (agreement, entropy).",
         "POSIX (Chatterjee 2024) · Sτ (Errica 2025)\nTVD consistency · semantic entropy\n(Kuhn 2023; Farquhar 2024)"),
        ("barsdown", "Benchmark robustness",
         "How much scores drop when\nbenchmark prompts are perturbed.",
         "Perturb prompts, re-score the model,\nreport the drop (model-level).",
         "BrittleBench (Romanou 2026)\nPromptEval (Polo 2024)"),
        ("pie", "Variance decomposition",
         "Which share of score variance\ncomes from which source.",
         "ANOVA / generalizability theory\nover repeated measurements.",
         "IR evaluation (Urbano 2013)\nLLM brand audits (Żatuchin 2026)\nρ_u embeddings (Cox 2025)"),
    ]
    for i, (icon, head, what, how, ex) in enumerate(cards):
        x = 0.41 + i * 4.23
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.50, 4.03, 4.15,
              fill=KIT_BLUE_TINT if i < 2 else KIT_GREEN_TINT,
              line=KIT_BLUE if i < 2 else KIT_GREEN)
        icon_station(s, x + 0.55, 1.68, icon, None)
        T(s, x + 1.05, 1.82, 2.9, 0.5, [head], size=16, bold=True)
        T(s, x + 0.18, 2.62, 3.7, 0.75,
          [[("What:  ", {"bold": True}), (what.replace("\n", " "), {})]],
          size=14)
        T(s, x + 0.18, 3.45, 3.7, 0.75,
          [[("How:  ", {"bold": True}), (how.replace("\n", " "), {})]],
          size=14)
        T(s, x + 0.18, 4.35, 3.7, 1.2, ex.split("\n"), size=12, color=MUTED)
    T(s, 0.41, 5.80, 12.5, 0.35,
      ["Missing in all three: the phrasing share of per-question task success, separated from sampling noise."],
      size=14, color=MUTED)
    chevron(s, "A rich zoo. Open question for Result 1: many measurements, or one measurement with many names?")


def s_functional_information(prs):
    s = kit_slide(prs, "Functional information", None,
                  notes="~90 s. Define FI first (Szostak/Hazen), then the two instantiations for prompts: "
                        "FI_in over the input space, FI_out over the output space. Derivation verified in "
                        "Section_7 notes: FI_in is the direct Hazen substitution; FI_out is the log N minus H "
                        "form, i.e. KL from uniform over the answer space. To our knowledge the first use of "
                        "FI to evaluate LLMs.")
    # definition
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 1.45, 12.5, 1.05,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 0.65, 1.58, 8.6, 0.5,
      [[("FI  =  −log₂ ( fraction of possibilities that achieve the function )",
         {"size": 17, "bold": True})]])
    T(s, 0.65, 2.08, 8.5, 0.35,
      [[("Rare success = many bits; common success = few bits.  (Szostak 2003; Hazen et al. 2007)",
         {"size": 12.5, "color": MUTED})]])
    for i, (h, b) in enumerate([("space", "all possibilities"),
                                ("test", "function ≥ threshold?"),
                                ("bits", "−log₂(m/N)")]):
        x = 9.35 + i * 1.22
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.62, 1.05, 0.68,
              fill=WHITE, line=KIT_GREEN)
        T(s, x, 1.68, 1.05, 0.55, [[(h, {"size": 11, "bold": True})],
                                   [(b, {"size": 9})]], align=PP_ALIGN.CENTER)
        if i:
            arrow(s, x - 0.15, 1.96, x, 1.96, width_pt=1.25)
    # FI_in panel
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 2.80, 6.12, 2.75,
          fill=WHITE, line=KIT_BLUE)
    T(s, 0.62, 2.92, 5.7, 0.4, [[("FI_in · input space", {"size": 16, "bold": True, "color": KIT_BLUE})]])
    T(s, 0.62, 3.38, 5.7, 1.0,
      [[("Possibility space:  ", {"bold": True}), ("the rephrasings of a question", {})],
       [("Function:  ", {"bold": True}), ("the answer is correct (≥ threshold k)", {})]],
      size=14)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.62, 4.35, 5.7, 0.62,
          fill=KIT_BLUE_TINT, line=None)
    T(s, 0.80, 4.44, 5.4, 0.45,
      [[("FI_in(q, k) = −log₂( N_k / |U_q| )", {"size": 15, "bold": True})]])
    T(s, 0.62, 5.05, 5.7, 0.4,
      ["how rare a working phrasing is: the Hazen formula, one substitution"],
      size=12, color=MUTED)
    # FI_out panel
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.79, 2.80, 6.12, 2.75,
          fill=WHITE, line=KIT_BLUE)
    T(s, 7.00, 2.92, 5.7, 0.4, [[("FI_out · output space", {"size": 16, "bold": True, "color": KIT_BLUE})]])
    T(s, 7.00, 3.38, 5.7, 1.0,
      [[("Possibility space:  ", {"bold": True}), ("the semantic answers a question admits", {})],
       [("Function:  ", {"bold": True}), ("how sharply one prompt narrows that space", {})]],
      size=14)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.00, 4.35, 5.7, 0.62,
          fill=KIT_BLUE_TINT, line=None)
    T(s, 7.18, 4.44, 5.4, 0.45,
      [[("FI_out(x) = log₂|A_q| − H_sem(Y | x)", {"size": 15, "bold": True})]])
    T(s, 7.00, 5.05, 5.7, 0.4,
      ["restrictiveness: distance from a uniform answer distribution, in bits"],
      size=12, color=MUTED)
    T(s, 0.41, 5.72, 12.5, 0.4,
      [[("To our knowledge, the first use of functional information to evaluate LLMs ",
         {"size": 14, "bold": True}),
        ("(the construction itself follows Hazen et al. 2007).", {"size": 14, "color": MUTED})]])
    chevron(s, "One ruler for everything that follows: −log₂(surviving fraction), only the possibility space changes.")


def s_ambigqa(prs):
    s = kit_slide(prs, "The testbed: AmbigQA", None,
                  notes="~60 s. Real Google queries via Natural Questions; annotators enumerated each "
                        "question’s valid readings, wrote one disambiguated rewrite per reading with "
                        "its own gold answers, and attached their evidence snippets. That gives us a "
                        "human-written specificity dial (FI_spec = log2 m0), the union-gold grading "
                        "protocol, and the reading-not-recall design. Example is illustrative.")
    # left: worked example
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 1.50, 6.15, 3.75,
          fill=WHITE, line=KIT_BLUE)
    T(s, 0.62, 1.60, 5.8, 0.35, ["Example: one question, two valid readings"],
      size=14, bold=True, color=KIT_BLUE)
    icon_bubble(s, 0.70, 2.02, 5.55, 0.60,
                "“When did the movie The Nun come out?”", size=13.5)
    arrow(s, 2.55, 2.72, 1.95, 3.02)
    arrow(s, 4.05, 2.72, 4.65, 3.02)
    for x, reading, ans in [(0.70, "the 2018 horror film", "gold answer: 2018"),
                            (3.50, "the 2013 French film", "gold answer: 2013")]:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.08, 2.75, 0.85,
              fill=KIT_BLUE_TINT, line=KIT_BLUE)
        T(s, x + 0.12, 3.16, 2.55, 0.35, [reading], size=13, bold=True)
        T(s, x + 0.12, 3.50, 2.55, 0.35, [ans], size=12.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.70, 4.15, 5.55, 0.85,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 0.85, 4.24, 5.3, 0.65,
      [[("FI_spec = log₂ m₀ = 1 bit", {"size": 14, "bold": True})],
       [("m₀ = number of valid readings (here 2)", {"size": 12, "color": MUTED})]])
    # right: what the dataset provides
    T(s, 6.90, 1.55, 6.0, 0.4, ["What AmbigQA provides (Min et al. 2020)"],
      size=16, bold=True, color=KIT_BLUE)
    bullets = [
        "real user questions (Google queries, via Natural Questions)",
        "annotators enumerate each question’s valid readings (m₀)",
        "one disambiguated rewrite + gold answers per reading",
        "the annotators’ own evidence snippets per question",
        "our sample: 2,002 validation questions, 150 after filters",
    ]
    y = 2.05
    for it in bullets:
        T(s, 7.04, y, 5.85, 0.55, [[("·  ", {"bold": True}), (it, {})]],
          size=14)
        y += 0.62
    # why it fits
    for i, (head, txt) in enumerate([
            ("Model-free dial", "both levels human-written:\nambiguous ↔ disambiguated"),
            ("Honest grading", "several valid answers →\nunion-gold protocol"),
            ("Reading, not recall", "shared evidence text\nfor every rephrasing")]):
        x = 0.41 + i * 4.23
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.45, 4.03, 0.95,
              fill=ROW_ALT, line=GRAY)
        T(s, x + 0.15, 5.53, 3.75, 0.35, [head], size=13.5, bold=True)
        T(s, x + 0.15, 5.86, 3.75, 0.5, txt.split("\n"), size=12, color=MUTED)
    chevron(s, "Ambiguity makes specificity measurable: FI_spec = log₂ m₀ bits, straight from annotations.",
            top=6.55, size=13.5)


def s_gap(prs):
    s = kit_slide(prs, "What is missing, and what we ask", None,
                  notes="~90 s. Three concrete gaps, then the research question they force. RQ1-RQ3 "
                        "return on the contributions slide; keep the numbering visible.")
    gaps = [
        ("scatter", "Two disconnected\nfamilies",
         "Indices either describe the answer distribution with no gold reference "
         "(POSIX, Sτ, semantic entropy, ρ_u), or reference task success without "
         "separating wording from sampling noise (ProSA, BrittleBench, Cao)."),
        ("pie", "No phrasing share\nof success",
         "Wording shares of an outcome exist: for sentiment (Żatuchin), answer "
         "embeddings (Cox), aggregate benchmark scores (Messing). For per-question "
         "task success under sampling: unmeasured."),
        ("dial", "No discriminant\nvalidation",
         "Metrics covary with training knobs (POSIX falls with few-shot exemplars), "
         "but none is validated by a planted dial with controls, and none is "
         "recomputed under a second paraphrase generator (PTEB: open)."),
    ]
    for i, (icon, head, txt) in enumerate(gaps):
        y = 1.42 + i * 1.12
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, y, 12.5, 1.00,
              fill=WHITE, line=KIT_RED)
        icon_station(s, 1.0, y + 0.12, icon, None)
        T(s, 1.75, y + 0.14, 3.1, 0.72, head.split("\n"), size=14.5, bold=True)
        T(s, 4.95, y + 0.10, 7.7, 0.85, [txt], size=12.5)
    arrow(s, 6.66, 4.70, 6.66, 4.98, color=KIT_GREEN, width_pt=2.5)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 5.04, 12.5, 0.92,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 0.62, 5.16, 12.1, 0.7,
      [[("Research question:  ", {"size": 15, "bold": True, "color": KIT_GREEN}),
        ("how can an LLM’s sensitivity to the phrasing of a question and to semantic change through added context be quantified, per question, in a way that is valid and actionable?", {"size": 15})]])
    T(s, 0.62, 6.22, 12.1, 0.4,
      [[("Answered through three contributions: ", {"size": 12.5, "color": MUTED, "bold": True}),
        ("a measurement model (C1), the missing metric ρ_F (C2), and a prompt "
         "checker (C3).", {"size": 12.5, "color": MUTED})]])
    s.notes_slide.notes_text_frame.text = (
        "~90 s. Gap wording is literature-proofed (verified 2026-08-10); if challenged: "
        "ProSA/BrittleBench/Cao DO score success per question but never separate wording from "
        "sampling noise (single greedy / T=0 / raw spread). Messing 2026 measures a wording "
        "share of MMLU answer-key success WITH noise separation, but at aggregate benchmark "
        "level; Zatuchin = sentiment, Cox = embeddings. POSIX is shown to fall under few-shot "
        "exemplars: training-knob covariation, not a planted discriminant dial. The "
        "paraphrase-generator-swap negative is verified clean (PTEB lists it as open future "
        "work). RQ mapping: phrasing = paraphrases; semantic change through added context = "
        "the specificity dial; valid = interventions; actionable = the checker.")


def s_cycle(prs):
    s = kit_slide(prs, "Research cycle (updated)", None,
                  notes="~60 s. Same skeleton as the January cycle; question and answer evolved with the "
                        "evidence. The January version asked for a bit-cost dose-response; a two-point design "
                        "cannot identify per-bit dosage, so the RQ became a measurement-model question.")
    box(s, 3.55, 1.50, 6.2, 1.45, [
        [("Practical Problem", {"bold": True, "size": 14})],
        [("Users cannot tell in advance whether phrasing will decide the outcome; "
          "suboptimal phrasing silently costs accuracy.", {"size": 12.5})],
    ], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 10.05, 2.55, 2.90, 2.75, [
        [("Research Question", {"bold": True, "size": 14})],
        [("How can sensitivity to phrasing and to semantic change through added "
          "context be quantified, per question, validly and actionably?", {"size": 12.5})],
    ], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 3.55, 4.95, 6.2, 1.55, [
        [("Research Problem", {"bold": True, "size": 14})],
        [("Indices describe the answer distribution or aggregate success; the "
          "noise-separated phrasing share of per-question success is unmeasured.", {"size": 12.5})],
    ], anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.38, 2.55, 2.90, 2.75, [
        [("Research Answer", {"bold": True, "size": 14})],
        [("A three-axis measurement model (accuracy, ρ_F, H_sem); each axis "
          "validated by its own intervention; plus a one-forward-pass prompt "
          "checker.", {"size": 12.5})],
    ], anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 9.75, 2.23, 10.55, 2.55)
    T(s, 10.75, 2.05, 2.2, 0.3, ["motivates"], size=12, color=MUTED)
    arrow(s, 10.55, 5.30, 9.75, 5.73)
    T(s, 10.60, 5.42, 2.6, 0.3, ["defines"], size=12, color=MUTED)
    arrow(s, 3.55, 5.73, 2.70, 5.30)
    T(s, 1.05, 5.52, 2.6, 0.3, ["provides solution to"], size=12, color=MUTED)
    arrow(s, 2.70, 2.55, 3.55, 2.23)
    T(s, 0.85, 1.86, 2.8, 0.3, ["contributes to the solution"], size=12,
      color=MUTED)
    T(s, 0.41, 6.62, 12.5, 0.35,
      [[("Updated from January:  ", {"bold": True, "size": 11.5, "color": MUTED}),
        ("the initial dose-response question is not identifiable from a two-point "
         "design (backup: pivots).", {"size": 11.5, "color": MUTED})]])


def s_design(prs):
    s = kit_slide(prs, "Design: one factorial, two planned interventions", None,
                  notes="~90 s. Walk the icons left to right: dataset, filters, two levels per question, "
                        "ten rephrasings, three models with ten samples, dual-gold scoring. Then the two "
                        "dials and the correctness contract.")
    stations = [
        (1.30, "can", "AmbigQA", "2,002 real user\nquestions"),
        (3.28, "funnel", "2 filters", "≥ 2 valid readings;\nevidence covers gold"),
        (5.26, "levels", "2 levels", "ambiguous L0 ↔\ndisambiguated L1"),
        (7.24, "gear", "× 10 rephrasings", "per question\nand level"),
        (9.22, "cubes", "3 models", "× 10 samples\neach"),
        (11.20, "check", "scoring", "against target and\nunion gold sets"),
    ]
    for i, (cx, kind, label, sub) in enumerate(stations):
        icon_station(s, cx, 1.62, kind, label, sub)
        if i:
            arrow(s, cx - 1.62, 2.05, cx - 1.02, 2.05)
    T(s, 2.15, 3.32, 2.3, 0.3, ["150 questions"], size=13, bold=True,
      color=KIT_BLUE, align=PP_ALIGN.CENTER)
    # dials
    for x, head, txt in [
            (0.41, "Dial 1 · Specificity (acts on the question)",
             "L0 ↔ L1;  FI_spec = log₂ m₀ bits, m₀ = number of annotator-listed valid readings"),
            (6.79, "Dial 2 · Generator width (acts on the rephrasings)",
             "narrow / production / wide / swap arms; identical filters; 50 questions, both levels")]:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.15, 6.12, 0.95,
              fill=KIT_GREEN_TINT, line=KIT_GREEN)
        icon_station(s, x + 0.55, 4.22, "dial", None)
        T(s, x + 1.15, 4.24, 4.85, 0.4, [head], size=14, bold=True)
        T(s, x + 1.15, 4.62, 4.85, 0.45, [txt], size=12)
    # held constant + scale
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 5.35, 9.2, 0.78,
          fill=WHITE, line=KIT_RED)
    icon_station(s, 0.95, 5.40, "lock", None)
    T(s, 1.55, 5.44, 8.0, 0.65,
      [[("Held constant:  ", {"bold": True}),
        ("identical evidence text everywhere; gold answers fixed across levels; "
         "the paraphrase generator is never an evaluated model.", {})]], size=14)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.81, 5.35, 3.1, 0.78,
          fill=ROW_ALT, line=GRAY)
    T(s, 9.97, 5.42, 2.85, 0.65,
      [[("≈ 90,000 graded answers", {"bold": True, "size": 13})],
       [("per gold set, on bwUniCluster", {"size": 11.5, "color": MUTED})]])
    chevron(s, "Only two things ever vary by design: the question’s specificity, and the width of its rephrasings.")


def s_instrument(prs):
    s = kit_slide(prs, "The instrument: a paraphrase universe", None,
                  notes="~60 s. Left to right: one question in, ten verified rephrasings out. The two gates "
                        "hold meaning and answerability constant, so whatever varies downstream is form. "
                        "Everything is seeded and cached; every rejection is logged per gate.")
    stations = [
        (1.30, "bubble", "1 question", "per specificity\nlevel"),
        (3.42, "gear", "Phi-4 drafts", "8 personas × 15;\nup to 120 candidates"),
        (5.54, "funnel", "Gate 1: meaning", "bidirectional NLI\nentailment ≥ 0.9"),
        (7.66, "funnel", "Gate 2: answerable", "gold-preservation\njudge"),
        (9.78, "dedup", "dedup + cap", "Levenshtein;\nmax 10"),
        (11.66, "grid", "universe U_q", "10 stimuli, never\nground truth"),
    ]
    for i, (cx, kind, label, sub) in enumerate(stations):
        icon_station(s, cx, 2.10, kind, label, sub)
        if i:
            arrow(s, cx - 1.72, 2.52, cx - 1.04, 2.52)
    T(s, 0.41, 4.30, 12.5, 0.4,
      ["Seeded and cached: byte-identical reruns. Every rejection is logged per gate; the gates themselves become a finding (Result 3)."],
      size=14, color=MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 4.90, 12.5, 0.85,
          fill=KIT_BLUE_TINT, line=KIT_BLUE)
    T(s, 0.62, 5.00, 12.1, 0.65,
      [[("Why gates, not free generation:  ", {"bold": True}),
        ("meaning and answerability are held constant by construction, so the universe varies "
         "only in form. FI_in’s possibility space is exactly this universe.", {})]], size=14)
    chevron(s, "Meaning held constant by construction: whatever varies downstream is form, not content.")


def s_axes(prs):
    s = kit_slide(prs, "The measurement model", None,
                  notes="~2 min. THE slide. One representative per axis; last column: each axis has its own "
                        "dial. The dial evidence is what the results section demonstrates.")
    data = [
        ["Axis", "Question it answers", "Representative", "Its own dial"],
        ["Competence", "How well does the\nmodel do?",
         "graded accuracy per question\n(mean over k = 10 samples), union gold",
         "Specificity: +6–13 pts\n(BH 3/3, Holm 2/3)"],
        ["Formulation\nsensitivity", "Does the wording\ndecide success?",
         "ρ_F ∈ [0,1]: share of success variance\nfrom phrasing (ICC; hierarchical estimator)",
         "Generator width:\nrises 3/3 (sig. in Qwen)"],
        ["Output\ndispersion", "How scattered are\nthe answers?",
         "H_sem (bits): entropy over\nsemantic answer clusters",
         "no dial of its own; weak\nspecificity response"],
        ["Manipulated\nvariable", "How underspecified\nis the question?",
         "FI_spec = log₂(m₀ / m_valid) bits;\nfrom annotations, model-free",
         "(this is Dial 1,\nnot an outcome)"],
    ]
    add_table(s, 0.41, 1.55, 12.5, data, col_widths=[2.2, 2.7, 4.8, 2.8],
              row_height=0.92, size=13, header_size=14, first_col_bold=True,
              highlight_rows={2})
    chevron(s, "One ruler: every quantity is −log₂(surviving fraction) over its own possibility space.",
            top=6.35)


def s_result_collapse(prs):
    s = kit_slide(prs, "Result 1: the zoo collapses into the three axes", None,
                  notes="~90 s. Every published index lands on one of the three axes, with graded "
                        "evidence: dispersion satellites by exact identity (proofs, max deviation "
                        "4.4e-16), competence satellites analytically / at rho = -.9997, sensitivity "
                        "satellites by empirical convergence (weaker, and labeled as such). POSIX is "
                        "the one index that cannot be assigned (non-discriminating). Horn confirms "
                        "no fourth dimension.")
    cols = [
        (0.41, KIT_BLUE, KIT_BLUE_TINT, "Competence", "accuracy F̄  (union gold)",
         [("AUFI, “ability in bits”",
           "≡ accuracy: analytic under binary scoring; ρ = −.9997 graded"),
          ("binary FI_in curve",
           "encodes a single number per cell")]),
        (4.64, KIT_GREEN, KIT_GREEN_TINT, "Formulation sensitivity", "ρ_F",
         [("ρ_u (Cox et al. 2025)",
           "same ICC estimand on embeddings, no gold; converges, ρ ≈ +.67"),
          ("performance spread (Cao et al. 2024)",
           "the uncorrected variant: no sampling-noise subtraction"),
          ("“islands of function” step size",
           "tracks ρ_F at +.83 to +.93")]),
        (8.87, KIT_ORANGE, RGBColor(0xFA, 0xF0, 0xDC), "Output dispersion",
         "H_sem  (= semantic entropy, adopted)",
         [("Sτ agreement (Errica et al. 2025)",
           "≡ H_sem / log₂|A|,  error ≤ 2×10⁻¹⁶"),
          ("FI_out(fixed)  ·  Var[FI_out]",
           "≡ log₂ m₀ − H_sem (exact)  ·  ≡ Var[H_sem]"),
          ("|A_q| · variation ratio · TVD consistency",
           "functions of the same answer clustering")]),
    ]
    for x, lc, fc, axis, rep, sats in cols:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.45, 4.03, 0.82,
              fill=fc, line=lc)
        T(s, x + 0.14, 1.52, 3.8, 0.35, [axis], size=14, bold=True, color=lc)
        T(s, x + 0.14, 1.86, 3.8, 0.35, [rep], size=13, bold=True)
        arrow(s, x + 2.0, 2.68, x + 2.0, 2.31, color=lc, width_pt=2.0)
        y = 2.74
        for name, rel in sats:
            h = 0.80
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 4.03, h,
                  fill=WHITE, line=GRAY)
            T(s, x + 0.13, y + 0.05, 3.8, 0.35, [name], size=12.5, bold=True)
            T(s, x + 0.13, y + 0.38, 3.8, 0.42, [rel], size=11.5, color=MUTED)
            y += h + 0.10
    T(s, 0.41, 5.48, 12.5, 0.3,
      [[("≡ exact identity (proved on every cell where defined)   ·   "
         "“converges / tracks” = empirical agreement, not an identity",
         {"size": 11.5, "color": MUTED})]])
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 5.80, 7.35, 0.68,
          fill=WHITE, line=GRAY)
    T(s, 0.56, 5.86, 7.1, 0.6,
      [[("POSIX ", {"bold": True, "size": 12}),
        ("is the one index that lands nowhere: it loads on ρ_F and H_sem "
         "equally (difference n.s., p = .42/.63/.054).", {"size": 12})]])
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.95, 5.80, 4.95, 0.68,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 8.10, 5.86, 4.7, 0.6,
      [[("No fourth axis: ", {"bold": True, "size": 12}),
        ("Horn retains exactly 3 factors (top 3: 74.6 %).", {"size": 12})]])
    chevron(s, "Every published index lands on one of the three axes; none adds a fourth.",
            top=6.55, size=13.5)


def s_result_specificity(prs):
    s = kit_slide(prs, "Result 2: the specificity dial", None,
                  notes="~2 min. Three panels, one message: the dial moves competence (union gold, the "
                        "honest score), leaves rho_F exactly flat (powered null), and moves dispersion "
                        "weakly. The union/target gap is the grading lottery, quantified in backup. "
                        "Question-clustered CIs and the full table: backup.")
    models = ["Qwen2.5-7B", "Llama-3.1-8B", "Mistral-7B"]
    # [SH] union +.064/+.125/+.119; target +.225/+.238/+.247
    add_chart(s, 0.41, 1.50, 4.25, 3.35, XL_CHART_TYPE.COLUMN_CLUSTERED,
              models,
              [("union gold (primary)", (0.064, 0.125, 0.119)),
               ("target gold", (0.225, 0.238, 0.247))],
              colors=[KIT_GREEN, GRAY], title="Δ accuracy (L1 − L0)",
              y_min=0, y_max=0.30, fmt="0.00")
    # [SH] H_sem deltas
    add_chart(s, 4.86, 1.50, 3.95, 3.35, XL_CHART_TYPE.COLUMN_CLUSTERED,
              models, [("Δ H_sem", (-0.124, -0.433, -0.139))],
              colors=[KIT_BLUE], title="Δ H_sem (bits)",
              y_min=-0.50, y_max=0.10, fmt="0.00", legend=False)
    # [SH] rho_F deltas, axis span matches panel A magnitude to show flatness
    add_chart(s, 9.01, 1.50, 3.90, 3.35, XL_CHART_TYPE.COLUMN_CLUSTERED,
              models, [("Δ ρ_F", (-0.002, 0.013, 0.013))],
              colors=[KIT_GREEN], title="Δ ρ_F (hierarchical)",
              y_min=-0.15, y_max=0.15, fmt="0.00", label_fmt="0.000",
              major_unit=0.05, legend=False)
    T(s, 0.41, 5.00, 12.5, 0.4,
      ["union gold = credit for any valid reading (the dataset’s own protocol); the union/target gap is the grading lottery: 47–72 % of the naive effect (backup)"],
      size=12.5, color=MUTED)
    chevron(s, "Accuracy rises +6–13 pts (BH 3/3, Holm 2/3; pooled p = 3×10⁻⁵).", top=5.55)
    chevron(s, "ρ_F does not move (n.s. 3/3, CIs exclude |Δ| > 0.04, both gold sets); H_sem falls weakly (Holm 1/3).",
            top=6.10)


def s_result_width(prs):
    s = kit_slide(prs, "Result 3: the generator-width dial", None,
                  notes="~2 min. Line charts: rho_F rises with width in all models (method-of-moments on "
                        "covered cells; per-arm hierarchical fit unidentifiable on the narrow arm, backup), "
                        "accuracy stays flat (Friedman >= .47; H_sem flat too). Swap strip: structure "
                        "survives an OLMo generator+judge swap. Together with Result 2: double dissociation.")
    arms = ["narrow", "production", "wide"]
    # [R6] MoM rho_F per arm
    add_chart(s, 0.41, 1.50, 6.0, 3.05, XL_CHART_TYPE.LINE,
              arms,
              [("Qwen2.5-7B", (0.356, 0.463, 0.519)),
               ("Llama-3.1-8B", (0.113, 0.135, 0.138)),
               ("Mistral-7B", (0.211, 0.230, 0.276))],
              colors=[C_QWEN, C_LLAMA, C_MISTRAL],
              title="ρ_F rises with width", y_min=0, y_max=0.6, fmt="0.0")
    # [R6] accuracy per arm (P3 rows)
    add_chart(s, 6.91, 1.50, 6.0, 3.05, XL_CHART_TYPE.LINE,
              arms,
              [("Qwen2.5-7B", (0.399, 0.418, 0.411)),
               ("Llama-3.1-8B", (0.310, 0.297, 0.303)),
               ("Mistral-7B", (0.366, 0.380, 0.374))],
              colors=[C_QWEN, C_LLAMA, C_MISTRAL],
              title="accuracy stays flat", y_min=0, y_max=0.5, fmt="0.0",
              major_unit=0.1)
    T(s, 0.41, 4.62, 12.5, 0.35,
      ["ρ_F: method of moments on covered cells; trend p = .004 / .079 / .051. Accuracy and H_sem: Friedman p ≥ .47 everywhere. Realized width 4.6 < 8.4 < 9.2 tokens; the NLI gate censors the wide arm (64 % rejections)."],
      size=12.5, color=MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 5.10, 12.5, 0.85,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 0.62, 5.20, 12.1, 0.65,
      [[("Swap ablation (first in this literature):  ", {"bold": True}),
        ("generator and judge replaced by OLMo-2-13B: per-cell ρ_F agreement +.41/+.27/+.51, "
         "consistent with attenuation by the split-half reliability (.38–.57); model ranking "
         "preserved. Levels are generator-relative, structure is generator-robust.", {})]], size=14)
    chevron(s, "Each axis moves under its own dial and only its own: the double dissociation is complete.",
            top=6.15)


def s_result_psychometrics(prs):
    s = kit_slide(prs, "Result 4: ρ_F as an instrument", None,
                  notes="~90 s. Chart: split-half reliability (population and ranking claims, not "
                        "per-question points) and out-of-sample payoff prediction per model. Right: "
                        "coverage and robustness in words.")
    models = ["Qwen2.5-7B", "Llama-3.1-8B", "Mistral-7B"]
    # [SH] reliability; [CV] payoff prediction (raw Spearman)
    add_chart(s, 0.41, 1.55, 6.1, 3.6, XL_CHART_TYPE.COLUMN_CLUSTERED,
              models,
              [("split-half reliability", (0.38, 0.52, 0.57)),
               ("payoff prediction (Spearman)", (0.61, 0.39, 0.70))],
              colors=[KIT_BLUE, KIT_GREEN], title=None, y_min=0, y_max=1.0,
              fmt="0.00")
    rows = [
        ("Coverage", "the classical estimator is undefined on 34–55 % of cells, exactly the "
                     "extremes; the hierarchical estimator covers 100 % with per-cell uncertainty."),
        ("Reliability", "split-half .38–.57: population-level and ranking claims are in scope, "
                        "per-question point claims are not."),
        ("Robustness", "ranking Qwen > Mistral > Llama survives both gold sets, both estimators, "
                       "both generators."),
        ("Utility", "ρ_F from the first 10 samples predicts the payoff of rephrasing on unseen "
                    "samples: “rephrase or give up”, decidable before trying."),
    ]
    y = 1.55
    for head, txt in rows:
        T(s, 6.85, y, 6.05, 0.3, [head], size=15, bold=True, color=KIT_BLUE)
        T(s, 6.85, y + 0.32, 6.05, 0.65, [txt], size=13)
        y += 1.03
    chevron(s, "An instrument with disclosed limits, and a decision-relevant payoff it predicts out of sample.",
            top=6.15)


def s_probe(prs):
    s = kit_slide(prs, "The prompt checker", None,
                  notes="~90 s. Top: how the probe works, one forward pass, no generation. Bottom left: "
                        "the two headline stats (in-domain and zero-shot, head vs text baseline). Right: "
                        "what it can and cannot do; the rho_F head at chance is an honest null.")
    stations = [
        (1.30, "bubble", "prompt", None),
        (3.60, "blockarrow", "one forward pass", "no generation"),
        (5.90, "slabs", "hidden state", "mid depth, last\nprompt token"),
        (8.20, "head", "linear head", "trained by question-\ngrouped CV"),
        (10.80, "warn", "“underspecified”", "warning before\nany answer"),
    ]
    for i, (cx, kind, label, sub) in enumerate(stations):
        icon_station(s, cx, 1.55, kind, label, sub, label_w=2.3)
        if i:
            arrow(s, cx - 1.60, 1.95, cx - 0.95, 1.95)
    # [PE] AUROC: head vs strongest text baseline
    add_chart(s, 0.41, 3.45, 6.3, 2.9, XL_CHART_TYPE.COLUMN_CLUSTERED,
              ["Qwen2.5-7B", "Llama-3.1-8B", "Mistral-7B"],
              [("head, in-domain", (0.874, 0.873, 0.873)),
               ("best text baseline, in-domain", (0.756, 0.756, 0.756)),
               ("head, zero-shot holdout", (0.678, 0.670, 0.678)),
               ("best frozen text baseline", (0.587, 0.587, 0.587))],
              colors=[KIT_GREEN, GRAY, KIT_BLUE, RGBColor(0xC9, 0xC9, 0xC9)],
              title="detecting underspecified questions (AUROC)",
              y_min=0.0, y_max=1.0, fmt="0.0", labels=False)
    T(s, 7.00, 3.55, 5.9, 2.4,
      [[("Zero-shot transfer is the claim:  ", {"bold": True, "size": 14}),
        ("with in-domain labels, text baselines catch up; the head’s value is "
         "label efficiency.", {"size": 14})],
       [("", {})],
       [("At the shipped threshold it flags 67–73 % of held-out questions at "
         "precision ≈ .66.", {"size": 14})],
       [("", {})],
       [("Honest null: ", {"bold": True, "size": 14}),
        ("ρ_F itself is not linearly readable from the prompt’s hidden "
         "state (chance-level probes).", {"size": 14})]])
    chevron(s, "A warning before the model answers, from one forward pass; transfers without target-domain labels.",
            top=6.42)


def s_contributions(prs):
    s = kit_slide(prs, "Contributions", None,
                  notes="~60 s. The three contributions together answer the RQ: C1 what to "
                        "quantify, C2 how to measure and validate it, C3 how to make it "
                        "actionable.")
    cards = [
        ("C1", "RQ1", "A measurement model that tidies the metric zoo",
         "Three questions, three representatives. Reduction proofs collapse five published indices "
         "into one dispersion quantity; POSIX shown non-discriminating; exactly 3 factors (Horn).",
         KIT_BLUE, KIT_BLUE_TINT),
        ("C2", "RQ2", "ρ_F: the missing axis, with an estimator that works",
         "Per-question, gold-referenced, noise-corrected share of task success from phrasing. "
         "Hierarchical estimator (100 % coverage), disclosed reliability, validated by its own dial "
         "and the literature’s first paraphraser-swap ablation.",
         KIT_GREEN, KIT_GREEN_TINT),
        ("C3", "RQ3", "A prompt-checker artifact",
         "“Your prompt is underspecified” from one forward pass: zero-shot .67–.68 where "
         "matched text baselines reach .55–.59; label efficiency, positioned against Kossen (2024) "
         "and Zhang (2025).",
         KIT_ORANGE, RGBColor(0xFA, 0xF0, 0xDC)),
    ]
    for i, (c, _rq, head, txt, lc, fc) in enumerate(cards):
        y = 1.50 + i * 1.62
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, y, 12.5, 1.46,
              fill=fc, line=lc)
        T(s, 0.62, y + 0.10, 0.9, 0.6, [[(c, {"size": 22, "bold": True, "color": lc})]])
        T(s, 1.55, y + 0.12, 11.0, 0.4, [head], size=15, bold=True)
        T(s, 1.55, y + 0.52, 11.2, 0.85, [txt], size=13)
    chevron(s, "One measurement model, one new metric, one usable artifact: together they answer the research question.",
            top=6.42)


def s_rq_answers(prs):
    s = kit_slide(prs, "The research question, answered", None,
                  notes="~45 s. The whole talk in one slide: the RQ on top, the three contributions "
                        "as its answer. C1: what to quantify. C2: how to measure and validate it. "
                        "C3: how to make it actionable.")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, 1.45, 12.5, 0.88,
          fill=KIT_GREEN_TINT, line=KIT_GREEN)
    T(s, 0.62, 1.54, 12.1, 0.7,
      [[("RQ:  ", {"size": 14.5, "bold": True, "color": KIT_GREEN}),
        ("how can an LLM’s sensitivity to the phrasing of a question and to semantic "
         "change through added context be quantified, per question, in a way that is "
         "valid and actionable?", {"size": 14.5})]])
    rows = [
        ("C1", "What is there to quantify?",
         "Three distinct axes, not one number. Most published indices are provably ONE of "
         "them (dispersion); axes are non-redundant within bounds (≤ |0.14|) and dissociate "
         "under intervention.", KIT_BLUE, KIT_BLUE_TINT),
        ("C2", "How to measure the phrasing share, and validate it?",
         "ρ_F with a hierarchical estimator: gold-referenced, noise-corrected, 100 % coverage; "
         "moves under its own dial, survives a generator swap, reliability disclosed (.38–.57).",
         KIT_GREEN, KIT_GREEN_TINT),
        ("C3", "Can it be anticipated before the model answers?",
         "Partly: underspecification and dispersion read from one forward pass (zero-shot "
         ".67–.68 vs .55–.59 text); ρ_F itself does not (honest null).",
         KIT_ORANGE, RGBColor(0xFA, 0xF0, 0xDC)),
    ]
    for i, (c, q, a, lc, fc) in enumerate(rows):
        y = 2.55 + i * 1.42
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.41, y, 3.45, 1.28,
              fill=fc, line=lc)
        T(s, 0.58, y + 0.08, 3.1, 0.4, [[(c, {"size": 16, "bold": True, "color": lc})]])
        T(s, 0.58, y + 0.44, 3.15, 0.8, [q], size=13, bold=True)
        arrow(s, 3.95, y + 0.64, 4.35, y + 0.64, color=lc, width_pt=2.25)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.42, y, 8.49, 1.28,
              fill=WHITE, line=lc)
        T(s, 4.62, y + 0.10, 8.1, 1.1, [a], size=13)


def s_limitations(prs):
    s = kit_slide(prs, "Limitations and outlook", None,
                  notes="~45 s. Left: what bounds the claims. Right: what the measurement view opens up.")
    T(s, 0.41, 1.50, 6.0, 0.4, ["Limitations"], size=18, bold=True,
      color=KIT_RED)
    lims = [
        "One dataset (AmbigQA, factoid QA); the probe holdout is the only transfer evidence",
        "Levels are generator-relative: bits depend on the paraphrase generator G (structure does not)",
        "Reliability .38–.57 caps per-question use; 150 questions support population claims",
        "k = 10 samples: H_sem and answer-set sizes are lower bounds",
        "The NLI gate censors realizable width (64 % rejections in the wide arm)",
        "Three 7–8B open-weight models; no frontier-scale replication",
    ]
    y = 2.05
    for it in lims:
        T(s, 0.55, y, 5.95, 0.62, [[("·  ", {"bold": True}), (it, {})]],
          size=14)
        y += 0.72
    T(s, 6.95, 1.50, 6.0, 0.4, ["Outlook"], size=18, bold=True,
      color=KIT_GREEN)
    outs = [
        "Measurement models as the audit layer for LLM evaluation, not more leaderboard indices",
        "ρ_F-style shares as monitored quality attributes of LLM services (prompt-QA gates)",
        "Reference-class engineering: choose the generator G that models YOUR user population",
        "Interventional validation (dials, swaps) as the standard for new metrics",
    ]
    y = 2.05
    for it in outs:
        T(s, 7.09, y, 5.85, 0.85, [[("·  ", {"bold": True}), (it, {})]],
          size=14)
        y += 0.98


def s_takeaways(prs):
    s = kit_slide(prs, "Key takeaways", None,
                  notes="~45 s + questions. Four sentences and the dial matrix; that is the talk.")
    msgs = [
        "Prompt sensitivity is three measurements, not one number.",
        "Most published indices are provably a single dispersion quantity; we prove it, not correlate it.",
        "The missing quantity, ρ_F, is measurable per question, noise-corrected, and each axis moves only under its own dial.",
        "A one-forward-pass checker warns about underspecified prompts before the model answers.",
    ]
    for i, m in enumerate(msgs):
        shape(s, MSO_SHAPE.CHEVRON, 0.41, 1.62 + i * 0.72, 0.30, 0.24,
              fill=KIT_GREEN, line=None)
        T(s, 0.85, 1.56 + i * 0.72, 12.0, 0.6, [m], size=16, bold=True)
    dd = [
        ["", "Competence", "ρ_F", "H_sem"],
        ["Specificity dial", "▲ +6–13 pts (BH 3/3, Holm 2/3)", "flat", "▼ weak (1/3)"],
        ["Width dial", "flat", "▲ up (3/3)", "flat"],
        ["Generator swap", "–", "structure preserved", "–"],
    ]
    add_table(s, 1.7, 4.75, 9.9, dd, col_widths=[2.3, 3.2, 2.4, 2.0],
              row_height=0.42, size=13, header_size=13, first_col_bold=True,
              header_fill=KIT_GREEN)
    T(s, 0.41, 6.65, 12.5, 0.4, ["Thank you. Questions?"], size=18, bold=True,
      color=KIT_BLUE)


def s_future_work(prs):
    s = kit_slide(prs, "Future work", None,
                  notes="~30 s or skip to questions. Left: concrete next experiments. Right: broader directions.")
    T(s, 0.41, 1.50, 6.0, 0.4, ["Concrete next steps"], size=18, bold=True,
      color=KIT_BLUE)
    left = [
        "Ask-an-LLM baseline for the prompt checker (one small cluster arm)",
        "Second dataset for external validity; cross-domain transfer of ρ_F",
        "Multi-level specificity for a real dose-response curve (the January idea, done right)",
        "Good–Turing correction for answer-space size at small k",
        "Per-question confidence intervals via exact ICC methods (Feldt / Urbano)",
    ]
    y = 2.05
    for it in left:
        T(s, 0.55, y, 5.95, 0.8, [[("·  ", {"bold": True}), (it, {})]],
          size=14)
        y += 0.85
    T(s, 6.95, 1.50, 6.0, 0.4, ["Broader directions"], size=18, bold=True,
      color=KIT_GREEN)
    right = [
        "Instruction-side sensitivity: the same model for system prompts and task templates",
        "Auto-rephrasing agents that consult ρ_F: rephrase or give up, decided per question",
        "Prompt-QA gates in production pipelines, with ρ_F monitored like latency or cost",
        "A public paraphrase-universe benchmark so metrics become comparable across papers",
    ]
    y = 2.05
    for it in right:
        T(s, 7.09, y, 5.85, 0.8, [[("·  ", {"bold": True}), (it, {})]],
          size=14)
        y += 0.85


# --------------------------------------------------------------------------- #
# backup                                                                      #
# --------------------------------------------------------------------------- #


def s_backup_endpoints(prs):
    s = kit_slide(prs, "Backup · Specificity: full endpoint table",
                  "Level 1 − level 0, paired on 150 questions per model; question-clustered 95 % CIs.",
                  notes=None)
    data = [
        ["Endpoint (L1 − L0)", "Qwen2.5-7B", "Llama-3.1-8B", "Mistral-7B"],
        ["Accuracy, union gold (primary)",
         "+0.064 [−0.000, +0.128] †", "+0.125 [+0.067, +0.186] ‡", "+0.119 [+0.056, +0.182] ‡"],
        ["Accuracy, target gold (protocol)", "+0.225 ‡", "+0.238 ‡", "+0.247 ‡"],
        ["of which grading lottery", "72 %", "47 %", "52 %"],
        ["H_sem (bits)", "−0.124 †", "−0.433 ‡", "−0.139 †"],
        ["ρ_F (hierarchical)", "−0.002 (n.s.)", "+0.013 (n.s.)", "+0.013 (n.s.)"],
    ]
    add_table(s, 0.41, 1.80, 12.5, data, col_widths=[4.3, 2.75, 2.75, 2.7],
              row_height=0.46, size=12, header_size=12, first_col_bold=True,
              highlight_rows={1})
    T(s, 0.41, 4.75, 12.5, 1.3, [
        [("† significant under Benjamini–Hochberg · ‡ additionally Holm-robust (family of 12 Wilcoxon tests) · "
          "grading lottery = share of the target-gold gain from guessing the annotator’s pinned reading "
          "(Δ_target = Δ_union + Δ_targeting)", {"size": 11.5, "color": MUTED})],
        [("Pooled one-experiment tests (deltas averaged over the 3 correlated models): "
          "Δaccuracy +0.103 (p = 3×10⁻⁵) · ΔH_sem −0.232 (p = 6×10⁻⁵).",
          {"size": 12.5})],
        [("ρ_F’s null is properly powered: n = 150 per model, CIs exclude |Δ| > 0.04, replicated under both gold sets.",
          {"size": 12.5})],
    ])


def s_backup_probe(prs):
    s = kit_slide(prs, "Backup · Prompt checker: full numbers", None, notes=None)
    data = [
        ["Head (target)", "In-domain AUROC (nested CV)", "Best text baseline", "Zero-shot annotator holdout"],
        ["Underspecification", ".873–.874 (3/3 models)", ".756 (length)", ".670–.678 vs .545–.587 frozen text"],
        ["Dispersion (H_sem)", ".67 / .76 / .77", "≤ .62", "not evaluated"],
        ["Reliability (rank corr.)", ".44 / .51 / .35", "not evaluated", "not evaluated"],
        ["Formulation sensitivity (ρ_F)", "chance (.50/.58/.51; n.s. after correction)", "not evaluated", "reported as a null result"],
    ]
    add_table(s, 0.41, 1.60, 12.5, data, col_widths=[3.0, 3.6, 2.2, 3.7],
              row_height=0.5, size=11.5, header_size=11.5, first_col_bold=True)
    T(s, 0.41, 4.35, 12.5, 1.0, [
        [("Operating point (threshold 0.65): ", {"bold": True, "size": 12.5}),
         ("flags 67–73 % of held-out questions, precision ≈ .66, recall .77–.83; usable as a ranking "
          "or soft gauge.", {"size": 12.5})],
        [("Controls: refitted permutation nulls at every layer, nested-CV layer selection, matched "
          "text-baseline protocols (refit in-domain, frozen for transfer).", {"size": 12.5})],
    ])


def s_backup_problems(prs):
    s = kit_slide(prs, "Backup · Problems we faced", None,
                  notes="Each problem produced a methods improvement; that is the story.")
    rows = [
        ("The grading lottery",
         "Scoring one pinned reading of an ambiguous question inflates the specificity effect: "
         "47–72 % of the naive +0.22–0.25 was protocol artifact. Found via the union-gold "
         "control arm; honest effect +0.06–0.13. Now a Methods feature."),
        ("Outcome-dependent missingness",
         "The classical ICC estimator is undefined exactly on all-right/all-wrong cells (34–55 %). "
         "Fix: hierarchical beta-binomial; plus a trap: a flat prior on the cell mean pulls degenerate "
         "cells to ρ = 1 (9.2×); caught in simulation, regression-tested."),
        ("Weak identifiability on the narrow arm",
         "Per-arm hierarchical fits inverted the ordering (median 7 paraphrases, mostly degenerate "
         "cells); diagnosed by simulation as a regime problem, not bias; reported MoM + count-matched "
         "checks. Single-seed subsampling was seed-lucky; 5-seed ranges."),
        ("The gate censors the dial",
         "The NLI equivalence gate rejects 64 % of the wide arm; the filter, not the generator, bounds "
         "realizable paraphrase width. Every NLI-filtered evaluation inherits this; per-gate rejection "
         "counts are persisted."),
        ("Engineering",
         "SQLite WAL over Lustre caused a cross-node lock storm on the cluster; fixed with per-chain "
         "caches. Content-hash caching made mixed A100/H100 windows harmless."),
    ]
    y = 1.50
    for h, txt in rows:
        T(s, 0.41, y, 12.5, 0.3, [h], size=13, bold=True, color=KIT_BLUE)
        T(s, 0.63, y + 0.28, 12.1, 0.62, [txt], size=11.5)
        y += 1.00
    chevron(s, "Every confound we found ourselves became a disclosed control.", top=6.50, size=13)


def s_backup_pivots(prs):
    s = kit_slide(prs, "Backup · Pivots since January", None, notes=None)
    data = [
        ["Was (January)", "Became (August)", "Why"],
        ["Bit-cost dose-response of specificity",
         "Two-point manipulated variable + a second dial (generator width)",
         "per-bit dosage not identifiable from a 2-point design"],
        ["“Three orthogonal axes”",
         "Bounded non-redundancy (≤ |0.14|; equivalence only to .33–.47) + experimental dissociation",
         "equivalence not establishable at n = 150; dissociation is stronger evidence"],
        ["AUFI, “ability in bits”",
         "Graded accuracy; FI curve kept as presentation",
         "AUFI ≡ accuracy (ρ = −.9997); bit value depends on an arbitrary cap"],
        ["Evidence dial (3rd manipulation)", "Withdrawn",
         "manipulates answerability, not context; range restriction"],
        ["“First application of FI to prompts”",
         "First use of FI to evaluate LLMs; construction follows Hazen",
         "Hazen et al. (2007) already scored language by receiver response"],
        ["Target-gold scoring", "Dual scoring; union gold primary",
         "AmbigQA’s own protocol scores all valid readings; pinned-reading scoring adds a 47–72 % lottery"],
        ["Probe as novel detector", "Zero-shot transfer + paraphrase-distribution heads",
         "scooped in parts by Kossen et al. (2024), Zhang et al. (2025); repositioned after verification"],
    ]
    add_table(s, 0.41, 1.55, 12.5, data, col_widths=[3.6, 4.6, 4.3],
              row_height=0.58, size=11, header_size=12, first_col_bold=False,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    T(s, 0.41, 6.42, 12.5, 0.4,
      [["All analysis decisions that changed a number are disclosed in a forking-paths appendix (12 forks)."]],
      size=11.5, color=MUTED)


def s_backup_lottery(prs):
    s = kit_slide(prs, "Backup · The grading lottery, quantified",
                  "Δ_target = Δ_union + Δ_targeting; paired on n = 150 questions; identical cached responses re-scored.",
                  notes=None)
    data = [
        ["Model", "Δ target gold", "=  Δ union gold (ability)", "+  Δ targeting (lottery)", "Lottery share"],
        ["Qwen2.5-7B", "+0.225", "+0.064  (p = .031)", "+0.161", "72 %"],
        ["Llama-3.1-8B", "+0.238", "+0.125  (p = 4×10⁻⁵)", "+0.113", "47 %"],
        ["Mistral-7B", "+0.247", "+0.119  (p = 3×10⁻⁴)", "+0.128", "52 %"],
    ]
    add_table(s, 0.41, 1.85, 12.5, data, col_widths=[2.4, 2.2, 3.5, 2.8, 1.9],
              row_height=0.5, size=12, header_size=12, first_col_bold=True)
    T(s, 0.41, 4.10, 12.5, 1.6, [
        [("Why it happens:  ", {"bold": True, "size": 12.5}),
         ("an ambiguous question has several valid readings; scoring only one pinned reading counts "
          "“guessed the annotator’s reading” as ability. Disambiguation then looks 2–4× "
          "better than it is.", {"size": 12.5})],
        [("Protocol consequence:  ", {"bold": True, "size": 12.5}),
         ("AmbigQA’s own evaluation scores against ALL valid readings; union gold is the licensed score "
          "and our primary endpoint.", {"size": 12.5})],
        [("Robustness:  ", {"bold": True, "size": 12.5}),
         ("ρ_F is gold-robust: per-cell agreement +.53/+.69/+.68 across scorings, identical ordering, "
          "identical dial-null.", {"size": 12.5})],
    ])
    chevron(s, "We found and removed our own confound: 47–72 % of the naive effect was the grading rule.",
            top=6.30, size=13)


def s_backup_stats(prs):
    s = kit_slide(prs, "Backup · Statistical hygiene",
                  "All inference numbers quoted from one artifact: data/stats_hygiene.md.", notes=None)
    rows = [
        ("Declared family",
         "12 paired Wilcoxon tests (4 endpoints × 3 models), Holm AND Benjamini–Hochberg; "
         "effect sizes with question-clustered bootstrap CIs (5,000 resamples). Disclosed failure: "
         "Qwen’s primary Δ_union survives BH (.045) but not Holm (.15)."),
        ("Models are correlated measurements, not replications",
         "Per-question accuracy deltas correlate at Spearman .52–.64 across models (H_sem: .37–.44). "
         "Pooled per-question tests: Δaccuracy +0.103 (p = 3.3×10⁻⁵), ΔH_sem "
         "−0.232 (p = 6.1×10⁻⁵)."),
        ("One test per measurement",
         "FI_out(fixed) is an affine relabeling of H_sem; its paired test IS the H_sem test and is "
         "excluded from the family."),
        ("Reliability policy",
         "Split-half on disjoint paraphrase halves (200 splits, Spearman–Brown): .38/.52/.57. The "
         "k = 20 arm contains the k = 10 sample; correlating a statistic with its superset measures "
         "overlap, not stability; retired."),
        ("Forking paths",
         "12 analysis decisions that changed a number are disclosed with the option not taken."),
    ]
    y = 1.70
    for head, txt in rows:
        T(s, 0.41, y, 12.5, 0.3, [head], size=13, bold=True, color=KIT_BLUE)
        T(s, 0.63, y + 0.28, 12.1, 0.62, [txt], size=11.5)
        y += 0.99


def s_backup_width(prs):
    s = kit_slide(prs, "Backup · Width dial: full numbers",
                  "4 arms: narrow (3 minimal personas, T=0.5) · production (8 personas, T=0.8) · wide (8 diverse personas, T=1.0) · swap (OLMo-2-13B).",
                  notes=None)
    data = [
        ["Quantity (prediction)", "Qwen2.5-7B", "Llama-3.1-8B", "Mistral-7B"],
        ["σ²_B rises (narrow→prod→wide)", ".0256→.0258→.0301  p=.054", ".0108→.0140→.0153  p=.021", ".0210→.0205→.0249  p=.060"],
        ["ρ_F (MoM) rises", ".356→.463→.519  p=.0035", ".113→.135→.138  p=.079", ".211→.230→.276  p=.051"],
        ["Accuracy flat (Friedman p)", ".68", ".53", ".54"],
        ["H_sem flat (Friedman p)", ".47", ".57", ".62"],
        ["Swap: per-cell ρ_F agreement", "+.41", "+.27", "+.51"],
    ]
    add_table(s, 0.41, 1.95, 12.5, data, col_widths=[4.0, 2.9, 2.9, 2.7],
              row_height=0.45, size=11.5, header_size=12, first_col_bold=True)
    T(s, 0.41, 4.90, 12.5, 1.4, [
        [("Manipulation check: ", {"bold": True, "size": 12}),
         ("realized width (mean pairwise token edit distance) 4.6 < 8.4 < 9.2; 66 % of cells fully "
          "ordered; narrow < wide at p = 4×10⁻¹⁸.", {"size": 12})],
        [("Gate censoring: ", {"bold": True, "size": 12}),
         ("the wide arm loses 64 % of candidates to the NLI gate; the narrow arm is shaped by the dedup "
          "gate (84 % rejections, 66 % of universes at the NLI fallback). The equivalence filter, not "
          "the prompt instructions, sets the upper end of realizable width.", {"size": 12})],
        [("Estimator note: ", {"bold": True, "size": 12}),
         ("per-arm hierarchical fits are not identifiable on the narrow arm; MoM on covered cells is "
          "reported, with count-matched 5-seed robustness (Qwen stable, p = .001–.013).", {"size": 12})],
    ])


def s_backup_related(prs):
    s = kit_slide(prs, "Backup · Positioning against the closest work", None,
                  notes=None)
    data = [
        ["Work / index", "Object measured", "Level", "Gold-ref.", "Noise-corr.", "Validated by its own dial"],
        ["POSIX (Chatterjee et al. 2024)", "log-prob shift across prompts", "prompt pair", "✗", "✗", "✗"],
        ["Sτ / TVD / |A| / Var[FI_out]", "output dispersion (≡ H_sem)", "question", "✗", "✗", "✗"],
        ["Semantic entropy (Kuhn 2023; Farquhar 2024)", "output dispersion", "prompt", "✗", "✗", "✗"],
        ["BrittleBench (Romanou et al. 2026)", "accuracy drop under perturbation", "model", "✓", "✗", "✗"],
        ["ρ_u (Cox et al. 2025)", "embedding ICC over rephrasings", "question", "✗", "✗ ¹", "✗"],
        ["G-theory (Urbano 2013; Żatuchin 2026)", "score variance shares", "system / corpus", "✓ / ✗", "✓", "✗"],
        ["ρ_F (this work)", "share of task success from phrasing", "question × model", "✓", "✓", "✓ (2 dials + swap)"],
    ]
    add_table(s, 0.41, 1.55, 12.5, data, col_widths=[3.6, 3.2, 1.7, 1.2, 1.3, 1.5],
              row_height=0.5, size=11, header_size=11.5, first_col_bold=True,
              highlight_rows={7})
    T(s, 0.41, 5.75, 12.5, 0.7,
      [["¹ ρ_u compares between-paraphrase variance to a 1/n baseline; it does not subtract the "
        "within-phrasing term. POSIX does fall under few-shot exemplars (Chatterjee 2024): a training "
        "knob, not a planted discriminant dial. Probes: Kossen et al. (2024) predict semantic entropy "
        "from hidden states (same token position); Zhang et al. (2025) detect ambiguity on AmbigQA with "
        "transfer. Our delta: paraphrase-distribution targets + zero-shot transfer evaluation."]],
      size=11.5, color=MUTED)


def s_backup_formulas(prs):
    s = kit_slide(prs, "Backup · The quantities, formally", None, notes=None)
    rows = [
        ("FI_in(q, k)  =  −log₂ ( N_k(q) / |U_q| )",
         "bits demanded of the phrasing: fraction of the paraphrase universe whose graded success "
         "reaches threshold k. AUFI (area under the curve) ≡ accuracy (ρ = −.9997); "
         "presentation only."),
        ("ρ_F  =  σ²_B / (σ²_B + σ²_W)   ∈ [0, 1]",
         "one-way ICC over the paraphrase universe: the share of success variance attributable to "
         "phrasing, corrected for the binomial sampling noise of k = 10 draws. Primary estimator: "
         "hierarchical beta-binomial with empirical-Bayes priors on both mean and correlation."),
        ("H_sem(x)  =  −Σ_c  p_c log₂ p_c",
         "semantic entropy over NLI-clustered answers, pooled per cell; FI_out(fixed) = log₂ m₀ "
         "− H̄_sem is a unit conversion of the same number."),
        ("FI_spec  =  log₂ ( m₀ / m_valid )",
         "the manipulated variable: bits of interpretation-narrowing supplied by the question itself; "
         "m₀ = annotator-listed valid readings, m_valid = readings still admissible. L0: m_valid = "
         "m₀ (0 bits); L1: m_valid = 1 (log₂ m₀ bits). Model-free."),
        ("Generator-relativity (FIᴳ)",
         "all bit values are relative to the paraphrase generator G’s proposal distribution; levels "
         "change under a different G (swap arm: means shift), structure (orderings, per-cell ranks) "
         "does not."),
    ]
    y = 1.55
    for f, txt in rows:
        T(s, 0.41, y, 12.5, 0.3, [f], size=14, bold=True, color=KIT_BLUE)
        T(s, 0.63, y + 0.30, 12.1, 0.6, [txt], size=11.5)
        y += 1.02


# --------------------------------------------------------------------------- #


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=str, default=OUT)
    args = ap.parse_args()

    configure_logging("make_ksri_deck")
    root = load_config().repo_root()
    seminar = root.parent.parent
    src = seminar / TEMPLATE
    dst = seminar / args.out
    if not src.exists():
        logger.error("template deck not found: {}", src)
        return 1
    if dst.exists():
        # NEVER silently clobber a possibly hand-edited deck (lesson of 2026-08-10:
        # a regeneration overwrote manual edits; only PowerPoint's in-memory copy
        # saved them). Same convention as make_kit_deck: back up first.
        from datetime import datetime
        bak = dst.with_name(dst.stem + "_pre-regen_"
                            + datetime.now().strftime("%Y%m%d-%H%M%S") + ".pptx")
        shutil.copy2(dst, bak)
        logger.warning("existing deck backed up -> {} (regeneration REPLACES all "
                       "manual edits; port them or work on a copy)", bak.name)
    shutil.copy2(src, dst)

    prs = Presentation(str(dst))
    removed = drop_slides_from(prs, 2)
    logger.info("kept title slide, removed {} slides", removed)

    s_title(prs)
    s_agenda(prs)
    s_motivation(prs)
    s_background(prs)
    s_functional_information(prs)
    s_ambigqa(prs)
    s_gap(prs)
    s_cycle(prs)
    divider(prs, "Methods")
    s_design(prs)
    s_instrument(prs)
    s_axes(prs)
    divider(prs, "Results")
    s_result_collapse(prs)
    s_result_specificity(prs)
    s_result_width(prs)
    s_result_psychometrics(prs)
    s_probe(prs)
    s_contributions(prs)
    s_rq_answers(prs)
    s_limitations(prs)
    s_takeaways(prs)
    s_future_work(prs)
    divider(prs, "Backup")
    s_backup_endpoints(prs)
    s_backup_probe(prs)
    s_backup_problems(prs)
    s_backup_pivots(prs)
    s_backup_lottery(prs)
    s_backup_stats(prs)
    s_backup_width(prs)
    s_backup_related(prs)
    s_backup_formulas(prs)

    prs.save(str(dst))
    logger.info("wrote {} ({} slides)", dst.name, len(prs.slides._sldIdLst))
    print(f"DONE {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
